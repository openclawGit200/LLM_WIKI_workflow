#!/usr/bin/env python3
"""
File Ingest - 將各種檔案轉為乾淨的 Markdown
支援：doc/docx/pdf/pptx/xlsx/csv/txt/html/md
"""

import sys
import json
import os
import re
import shutil
import subprocess
import tempfile
import argparse
from datetime import datetime
from pathlib import Path

# ── 工具函式 ────────────────────────────────────────────────────

def run(cmd: list, capture=True) -> str:
    """執行 shell 命令，失敗則拋例外"""
    try:
        result = subprocess.run(
            cmd, capture_output=capture, text=True,
            encoding='utf-8', errors='replace', timeout=120
        )
        if result.returncode != 0:
            raise RuntimeError(f"Command failed: {' '.join(cmd)}\n{result.stderr}")
        return result.stdout
    except FileNotFoundError as e:
        raise RuntimeError(f"Tool not found: {cmd[0]}") from e


def clean_markdown(text: str) -> str:
    """清洗 markdown：移除過多空行、異常空白、腳本殘留"""
    # 移除 ANSI escape
    text = re.sub(r'\x1b\[[0-9;]*[a-zA-Z]', '', text)
    # 移除 Windows 換行
    text = text.replace('\r\n', '\n')
    # 移除 Word OLE 物件殘留（markitdown 常見）
    text = re.sub(r'EMBED\s*\*\s*MERGEFORMAT\s*', '', text)
    text = re.sub(r'\{\\[^}]*\}', '', text)   # Word 域殘留
    # 合併連續空行（超過2個換行 → 2個）
    text = re.sub(r'\n{3,}', '\n\n', text)
    # 移除行尾空白
    lines = [line.rstrip() for line in text.split('\n')]
    text = '\n'.join(lines)
    # 去除頭尾空白
    text = text.strip()
    return text


def slugify(name: str) -> str:
    """把檔案名稱轉成安全的 slug"""
    base = os.path.basename(name)
    name_without_ext = os.path.splitext(base)[0]
    # 移除不安全字元
    slug = re.sub(r'[^\w\s\-_]', '_', name_without_ext)
    slug = re.sub(r'[\s\-_]+', '_', slug)
    slug = slug.strip('_')
    return slug[:80] or 'untitled'


# ── 各類型轉換器 ────────────────────────────────────────────────

def convert_doc(path: str) -> str:
    """
    doc/docx/rtf/html → 乾淨 Markdown（via markitdown）
    - .doc/.rtf：先用 textutil 轉成 .docx，再由 markitdown 處理
    - .docx/.html/.htm：直接 markitdown
    - .txt：直接讀取（見 convert_txt）
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == '.txt':
        return convert_txt(path)

    if ext in ('.doc', '.rtf'):
        # textutil -outdir 有 Unicode 問題，改用 stdout + 手動寫檔
        tmp_dir = tempfile.mkdtemp()
        clean_name = 'converted_input'   # 避免 Unicode 檔名問題
        tmp_docx = os.path.join(tmp_dir, clean_name + '.docx')
        tmp_src  = os.path.join(tmp_dir, clean_name + ext)
        shutil.copy2(path, tmp_src)
        result = subprocess.run(
            ['textutil', '-convert', 'docx', '-stdout', tmp_src],
            capture_output=True, timeout=60
        )
        if result.returncode != 0:
            raise RuntimeError(f"textutil 轉換失敗：{result.stderr.decode(errors='replace')}")
        with open(tmp_docx, 'wb') as f:
            f.write(result.stdout)
        # 交給 markitdown 處理
        cmd = ['markitdown', tmp_docx]
        text = run(cmd)
        shutil.rmtree(tmp_dir, ignore_errors=True)
        return text
    elif ext in ('.html', '.htm', '.docx', '.pptx', '.xlsx', '.epub', '.odt'):
        cmd = ['markitdown', path]
        text = run(cmd)
        return text
    else:
        raise ValueError(f"convert_doc 不支援 {ext}")


def convert_pdf(path: str) -> str:
    """PDF → 純文字（via pdftotext）"""
    cmd = ['pdftotext', '-layout', path, '-']
    text = run(cmd)
    if not text.strip():
        cmd = ['pdftotext', path, '-']
        text = run(cmd)
    return text


def convert_pptx(path: str) -> str:
    """PPTX → Markdown（保留投影片分頁結構）"""
    from pptx import Presentation
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        slide_texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        slide_texts.append(text)
        if slide_texts:
            slides.append(f"## Slide {i}\n\n" + '\n'.join(f"- {t}" for t in slide_texts))
    return '\n\n'.join(slides) if slides else ""


def convert_xlsx(path: str) -> str:
    """XLSX → Markdown 表格"""
    try:
        import openpyxl
    except ImportError:
        raise RuntimeError("需先安裝 openpyxl: pip3 install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets_out = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        table_rows = []
        for row in rows:
            cells = [str(c) if c is not None else '' for c in row]
            table_rows.append(cells)
        md_lines = [f"### {sheet_name}\n",
                    "| " + " | ".join(table_rows[0]) + " |",
                    "| " + " | ".join(["---"] * len(table_rows[0])) + " |"]
        for row in table_rows[1:]:
            md_lines.append("| " + " | ".join(row) + " |")
        sheets_out.append('\n'.join(md_lines))
    return '\n\n'.join(sheets_out)


def convert_csv(path: str) -> str:
    """CSV → Markdown 表格"""
    import csv
    with open(path, encoding='utf-8-sig', errors='replace') as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return ""
    md_lines = [
        "| " + " | ".join(rows[0]) + " |",
        "| " + " | ".join(["---"] * len(rows[0])) + " |",
    ]
    for row in rows[1:]:
        md_lines.append("| " + " | ".join(row) + " |")
    return '\n'.join(md_lines)


def convert_txt(path: str) -> str:
    """純文字檔直接讀"""
    with open(path, encoding='utf-8-sig', errors='replace') as f:
        return f.read()


# ── 主轉換分派 ─────────────────────────────────────────────────

def convert_file(input_path: str) -> dict:
    """
    將檔案轉為 Markdown，回傳 metadata + 內容。
    """
    input_path = os.path.expanduser(input_path)
    if not os.path.exists(input_path):
        raise FileNotFoundError(f"檔案不存在：{input_path}")

    ext = os.path.splitext(input_path)[1].lower()
    file_size = os.path.getsize(input_path)
    file_mtime = datetime.fromtimestamp(os.path.getmtime(input_path)).isoformat()

    # 分派轉換器
    if ext in ('.doc', '.docx', '.rtf', '.html', '.htm'):
        text = convert_doc(input_path)
    elif ext == '.pdf':
        text = convert_pdf(input_path)
    elif ext == '.pptx':
        text = convert_pptx(input_path)
    elif ext == '.xlsx':
        text = convert_xlsx(input_path)
    elif ext == '.csv':
        text = convert_csv(input_path)
    elif ext in ('.txt', '.md', '.markdown'):
        text = convert_txt(input_path)
    else:
        raise ValueError(f"不支援的檔案類型：{ext}")

    # 清洗
    markdown = clean_markdown(text)

    if len(markdown) < 50:
        raise RuntimeError(f"轉換後內容少於50字，可能轉換失敗：{markdown[:200]}")

    return {
        "original_path": input_path,
        "original_name": os.path.basename(input_path),
        "file_type": ext,
        "file_size_bytes": file_size,
        "file_mtime": file_mtime,
        "markdown": markdown,
        "char_count": len(markdown),
    }


# ── CLI 入口 ───────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="File Ingest - 轉檔為 Markdown")
    parser.add_argument("file", help="輸入檔案路徑")
    parser.add_argument("--output-dir", default=None, help="輸出目錄（預設為 raw md/）")
    parser.add_argument("--vault-path", default=None, help="Obsidian vault 路徑")
    parser.add_argument("--frontmatter", action='store_true', default=True,
                        help="加入 YAML frontmatter（預設開啟）")
    parser.add_argument("--no-frontmatter", dest='frontmatter', action='store_false',
                        help="不加入 YAML frontmatter")
    parser.add_argument("--dry-run", action='store_true', help="僅顯示結果，不寫入檔案")
    args = parser.parse_args()

    # 決定輸出目錄
    if args.output_dir:
        out_dir = os.path.expanduser(args.output_dir)
    elif args.vault_path:
        out_dir = os.path.join(os.path.expanduser(args.vault_path), "raw md")
    else:
        import json as _json
        vault_cfg = os.path.expanduser("~/Library/Application Support/obsidian/obsidian.json")
        if os.path.exists(vault_cfg):
            with open(vault_cfg) as f:
                vaults = _json.load(f).get('vaults', {})
            open_v = [(n, v) for n, v in vaults.items() if v.get('open')]
            if open_v:
                out_dir = os.path.join(open_v[0][1]['path'], "raw md")
            else:
                out_dir = None
        else:
            out_dir = None

    if not out_dir:
        raise RuntimeError("無法自動判斷 vault 路徑，請用 --vault-path 指定")

    # 轉換
    result = convert_file(args.file)
    slug = slugify(result['original_name'])
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_filename = f"{slug}_{timestamp}.md"
    output_path = os.path.join(out_dir, output_filename)

    # 建目錄
    os.makedirs(out_dir, exist_ok=True)

    # 組合內容
    if args.frontmatter:
        fm = [
            "---",
            f"title: \"{result['original_name']}\"",
            f"source_file: \"{result['original_path']}\"",
            f"source_type: \"{result['file_type']}\"",
            f"source_size: {result['file_size_bytes']}",
            f"converted_at: \"{datetime.now().isoformat()}\"",
            f"char_count: {result['char_count']}",
            f"file_mtime: \"{result['file_mtime']}\"",
            "tags: []",
            "---",
            "",
        ]
        content = '\n'.join(fm) + '\n' + result['markdown']
    else:
        content = result['markdown']

    result['output_path'] = output_path
    result['output_dir'] = out_dir

    if args.dry_run:
        result['content_preview'] = content[:500]
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(content)
        result['markdown'] = None
        result['content_preview'] = content[:300]
        print(json.dumps(result, ensure_ascii=False, indent=2))
        print(f"\n✅ 已寫入：{output_path}", file=sys.stderr)


if __name__ == '__main__':
    main()
